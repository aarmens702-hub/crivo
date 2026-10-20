"""Messy Excel structure intelligence (capability roadmap B2.1).

Real workbooks are not tidy: a title row, blank rows, notes, several tables on
one sheet. pandas.read_excel assumes row 1 is the header and one table per
sheet, so it silently mangles these. This module reads the raw cell grid with
openpyxl (already a dependency, no new one) and reports, per sheet, where the
header really is and whether more than one table lives there, as findings a
person confirms. Keyless and pure (openpyxl + stdlib)."""

import openpyxl
import pytest

from crivo import excel_structure


def _wb(tmp_path, sheets):
    """sheets: {name: list-of-rows}; write a real .xlsx and return its path."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(list(row))
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def _one(findings, sheet):
    return next(f for f in findings if f["sheet"] == sheet)


def test_tidy_sheet_reports_header_row_1_and_clean(tmp_path):
    path = _wb(tmp_path, {"tidy": [["id", "amount"], [1, 10], [2, 20]]})
    f = _one(excel_structure.analyze_workbook(path), "tidy")
    assert f["header_row"] == 1
    assert f["kind"] == "clean"
    assert f["tables"] == 1


def test_title_and_notes_rows_push_the_header_down(tmp_path):
    path = _wb(
        tmp_path,
        {
            "report": [
                ["Quarterly Sales Report"],  # title, one cell
                [],  # blank
                ["region", "units", "revenue"],  # the real header, row 3
                ["west", 10, 100],
                ["east", 8, 80],
            ]
        },
    )
    f = _one(excel_structure.analyze_workbook(path), "report")
    assert f["header_row"] == 3
    assert f["preamble_rows"] == 2
    assert f["kind"] == "header-below-row-1"
    assert "3" in f["evidence"]


def test_two_tables_on_one_sheet_are_detected(tmp_path):
    path = _wb(
        tmp_path,
        {
            "twotab": [
                ["a", "b"],
                [1, 2],
                [3, 4],
                [],  # blank separator
                [],
                ["x", "y", "z"],
                [5, 6, 7],
            ]
        },
    )
    f = _one(excel_structure.analyze_workbook(path), "twotab")
    assert f["tables"] == 2
    assert f["kind"] == "multiple-tables"


def test_empty_sheet_is_reported_not_crashed(tmp_path):
    path = _wb(tmp_path, {"blank": [[], []]})
    f = _one(excel_structure.analyze_workbook(path), "blank")
    assert f["kind"] == "empty"
    assert f["header_row"] is None


def test_every_sheet_gets_exactly_one_finding(tmp_path):
    path = _wb(
        tmp_path,
        {
            "s1": [["h"], [1]],
            "s2": [["a", "b"], [1, 2]],
        },
    )
    findings = excel_structure.analyze_workbook(path)
    assert {f["sheet"] for f in findings} == {"s1", "s2"}
    assert len(findings) == 2


def test_findings_are_gradeable_and_suggest_a_fix(tmp_path):
    path = _wb(
        tmp_path,
        {"r": [["Title"], [], ["col_a", "col_b"], [1, 2]]},
    )
    f = _one(excel_structure.analyze_workbook(path), "r")
    assert f["grade"] in {"GATE", "HUMAN"}  # restructuring is a judgement call
    assert f["suggestion"]  # a concrete next step, e.g. skip N rows


_SALES_HEADER = ["region", "units", "revenue"]
_SALES_ROWS = [["west", 10, 100], ["east", 8, 80]]


def test_a_title_directly_above_the_header_is_preamble(tmp_path):
    """Regression for #105. A one-cell title with no blank row under it was
    joined to the table below and read as its header, so the sheet came back
    clean with the header on row 1. It must read the way the same table under
    a blank row does: header on row 2, one preamble row, GATE, with the same
    evidence and suggestion."""
    path = _wb(
        tmp_path,
        {
            "titled": [["Quarterly Sales Report"], _SALES_HEADER, *_SALES_ROWS],
            # the same table under a blank row: today's header-below-row-1
            # finding, the reference for the evidence and suggestion wording
            "blank_top": [[], _SALES_HEADER, *_SALES_ROWS],
        },
    )
    findings = excel_structure.analyze_workbook(path)
    f, ref = _one(findings, "titled"), _one(findings, "blank_top")
    assert f["header_row"] == 2
    assert f["preamble_rows"] == 1
    assert f["tables"] == 1
    assert f["kind"] == "header-below-row-1"
    assert f["grade"] == "GATE"
    assert f["evidence"] == ref["evidence"]
    assert f["suggestion"] == ref["suggestion"]  # skip the first 1 row


def test_a_merged_title_directly_above_the_header_is_preamble(tmp_path):
    """A title merged across the table's width (A1:C1) holds its text in A1
    only, so it is still one cell wide and still preamble."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "merged"
    ws.append(["Quarterly Sales Report"])
    ws.merge_cells("A1:C1")
    for row in [_SALES_HEADER, *_SALES_ROWS]:
        ws.append(row)
    path = tmp_path / "book.xlsx"
    wb.save(path)

    f = _one(excel_structure.analyze_workbook(path), "merged")
    assert f["header_row"] == 2
    assert f["preamble_rows"] == 1
    assert f["tables"] == 1
    assert f["kind"] == "header-below-row-1"
    assert f["grade"] == "GATE"


@pytest.mark.parametrize(
    "rows, header_row, preamble_rows, tables, kind, grade",
    [
        pytest.param(
            [
                ["Quarterly Sales Report"],
                ["As of 30 Sep 2026"],
                _SALES_HEADER,
                *_SALES_ROWS,
            ],
            3,
            2,
            1,
            "header-below-row-1",
            "GATE",
            id="title-and-as-of-line",
        ),
        pytest.param(
            [[None, "Quarterly Sales Report", None], _SALES_HEADER, *_SALES_ROWS],
            2,
            1,
            1,
            "header-below-row-1",
            "GATE",
            id="title-centred-in-column-b",
        ),
        pytest.param(
            [
                ["Quarterly Sales Report"],
                [],
                ["Region detail"],
                _SALES_HEADER,
                *_SALES_ROWS,
            ],
            4,
            3,
            1,
            "header-below-row-1",
            "GATE",
            id="title-blank-then-subtitle-on-the-header",
        ),
        pytest.param(
            [
                ["Quarterly Sales Report"],
                _SALES_HEADER,
                ["west", 10, 100],
                ["see note 3"],  # a narrow row inside the table is not preamble
                ["east", 8, 80],
            ],
            2,
            1,
            1,
            "header-below-row-1",
            "GATE",
            id="title-plus-a-narrow-row-mid-table",
        ),
        pytest.param(
            [
                ["Quarterly Sales Report"],
                ["a", "b"],
                [1, 2],
                [],
                ["x", "y", "z"],
                [5, 6, 7],
            ],
            2,
            1,
            2,
            "multiple-tables",
            "HUMAN",
            id="title-above-the-first-of-two-tables",
        ),
    ],
)
def test_only_the_narrow_rows_leading_into_the_header_are_preamble(
    tmp_path, rows, header_row, preamble_rows, tables, kind, grade
):
    """Every narrow row that leads straight into the first wide row of the
    first table is preamble; a narrow row after the header is part of the
    table, and the multiple-tables finding still names the real header."""
    path = _wb(tmp_path, {"s": rows})
    f = _one(excel_structure.analyze_workbook(path), "s")
    assert f["header_row"] == header_row
    assert f["preamble_rows"] == preamble_rows
    assert f["tables"] == tables
    assert f["kind"] == kind
    assert f["grade"] == grade


def test_a_header_with_one_label_reads_as_a_title(tmp_path):
    """The decided trade-off on #105: the width rule cannot tell a header with
    a single labelled cell from a title, so it reads as a title and the sheet
    goes to GATE, header on row 2. A person confirms a GATE before any skip;
    a wrong "clean" would read the title as the header with no one asked."""
    path = _wb(tmp_path, {"s": [["region", None, None], *_SALES_ROWS]})
    f = _one(excel_structure.analyze_workbook(path), "s")
    assert f["header_row"] == 2
    assert f["preamble_rows"] == 1
    assert f["kind"] == "header-below-row-1"
    assert f["grade"] == "GATE"


def test_a_narrow_data_row_under_a_full_header_stays_clean(tmp_path):
    """Pin beside #105: a narrow row is preamble only when it leads into the
    header; a data row with a missing value under a full header is data."""
    path = _wb(tmp_path, {"s": [["id", "amount"], ["007", None], ["012", 20]]})
    f = _one(excel_structure.analyze_workbook(path), "s")
    assert f["header_row"] == 1
    assert f["preamble_rows"] == 0
    assert f["kind"] == "clean"
    assert f["grade"] == "AUTO"


@pytest.mark.parametrize(
    "header",
    [
        pytest.param(["region", None, "revenue"], id="empty-cell"),
        pytest.param(["region", "   ", "revenue"], id="whitespace-cell"),
    ],
)
def test_a_header_with_one_blank_cell_stays_clean(tmp_path, header):
    """Pin beside #105: a header missing one label is still two or more wide,
    so it is the header, not a title."""
    path = _wb(tmp_path, {"s": [header, *_SALES_ROWS]})
    f = _one(excel_structure.analyze_workbook(path), "s")
    assert f["header_row"] == 1
    assert f["preamble_rows"] == 0
    assert f["kind"] == "clean"
    assert f["grade"] == "AUTO"


@pytest.mark.parametrize(
    "rows",
    [
        pytest.param(
            [_SALES_HEADER, ["North"], *_SALES_ROWS], id="group-label-under-header"
        ),
        pytest.param(
            [_SALES_HEADER, ["west", 10, 100], ["see note 3"], ["east", 8, 80]],
            id="note-mid-table",
        ),
    ],
)
def test_a_one_value_row_inside_a_table_does_not_move_the_header(tmp_path, rows):
    """Pin beside #105: a one-value row after the header neither splits the
    table nor moves the header."""
    path = _wb(tmp_path, {"s": rows})
    f = _one(excel_structure.analyze_workbook(path), "s")
    assert f["header_row"] == 1
    assert f["preamble_rows"] == 0
    assert f["tables"] == 1
    assert f["kind"] == "clean"
